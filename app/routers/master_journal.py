from __future__ import annotations

import json
import math
import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook, Workbook
from sqlalchemy import or_, and_, cast, Numeric, func
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Property, LocalGovernment, JournalField, JournalValue

router = APIRouter(prefix="/master-journal", tags=["Master Journal"])

ALIASES = {
    "property_no": {"assessment no"},
    "owner_name": {"name of occupier"},
    "lg_code": {"lg code"},
    "address": {"address"},
    "rating_area": {"rating area"},
}


def clean_label(value: Any) -> str:
    if value is None:
        return "unnamed_column"
    text = str(value).strip()
    return text if text else "unnamed_column"


def normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", clean_label(value).lower()).strip("_") or "unnamed_column"


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()
    text = str(value).strip()
    return text or None


def numeric_value(value: Any):
    text = clean_text(value)
    if text is None:
        return None
    candidate = text.replace(",", "").replace("₦", "").replace("N", "").replace("₦", "").strip()
    if candidate.startswith("(") and candidate.endswith(")"):
        candidate = "-" + candidate[1:-1]
    try:
        result = Decimal(candidate)
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def column_map(labels: list[str]) -> tuple[dict[str, int], list[dict[str, Any]]]:
    """Keep duplicate headings separate by occurrence, e.g. arrears and arrears__2."""
    counts: dict[str, int] = {}
    mapped: dict[str, int] = {}
    metadata = []
    for position, raw_label in enumerate(labels):
        label = clean_label(raw_label)
        base = normalize(label)
        counts[base] = counts.get(base, 0) + 1
        occurrence = counts[base]
        key = base if occurrence == 1 else f"{base}__{occurrence}"
        mapped[key] = position
        metadata.append({"field_key": key, "label": label, "normalized_label": base, "position": position, "occurrence": occurrence})
    return mapped, metadata


def resolve_standard_columns(labels: list[str]) -> dict[str, int]:
    resolved = {}
    for idx, label in enumerate(labels):
        norm = normalize(label)
        raw = clean_label(label).lower().strip()
        for standard, aliases in ALIASES.items():
            if standard not in resolved and (norm in {normalize(x) for x in aliases} or raw in aliases):
                resolved[standard] = idx
    return resolved


def safe_record(prop: Property) -> dict[str, Any]:
    return {
        "id": prop.id, "property_no": prop.property_no, "lg_code": prop.lg_code,
        "owner_name": prop.owner_name, "address": prop.address,
        "rating_area": prop.rating_area, "annual_value": prop.annual_value,
        "rate_due": prop.rate_due, "year": prop.year, "status": prop.status,
        "lg_id": prop.lg_id, "lg_prefix": prop.lg_prefix,
    }


def make_json_safe(value: Any) -> Any:
    """Convert missing and non-finite values to JSON-safe values."""
    if value is None:
        return None

    if isinstance(value, float):
        return value if math.isfinite(value) else None

    if isinstance(value, Decimal):
        return float(value) if value.is_finite() else None

    if isinstance(value, dict):
        return {
            str(key): make_json_safe(item)
            for key, item in value.items()
        }

    if isinstance(value, (list, tuple)):
        return [make_json_safe(item) for item in value]

    # Handle NumPy scalar values if any enter the response.
    if type(value).__module__.startswith("numpy") and hasattr(value, "item"):
        return make_json_safe(value.item())

    return value



@router.post("/worksheets")
async def get_worksheets(file: UploadFile = File(...)):
    contents = await file.read()
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        return {"worksheets": names}
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read this Excel workbook: {exc}")


@router.post("/import")
async def import_journal_excel(
    file: UploadFile = File(...),
    sheet_name: str = Form(...),
    local_government: str = Form(""),
    prefix: str = Form(""),
    db: Session = Depends(get_db),
):
    contents = await file.read()
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        
        selected_sheet = next(
            (
                name for name in workbook.sheetnames
                if name.strip() == sheet_name.strip()
            ),
            None,
        )

        if selected_sheet is None:
            raise HTTPException(
                status_code=400,
                detail="The selected worksheet does not exist in this workbook.",
            )

        sheet = workbook[selected_sheet]

        rows_iter = sheet.iter_rows(values_only=True)

        header_row = next(rows_iter, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="The selected worksheet is empty.")
        labels = [clean_label(v) for v in header_row]
        # Trim trailing empty Excel columns without changing internal blank headings.
        while labels and labels[-1] == "unnamed_column" and (header_row[len(labels)-1] is None or str(header_row[len(labels)-1]).strip() == ""):
            labels.pop()
        standard = resolve_standard_columns(labels)
        if "property_no" not in standard or "owner_name" not in standard:
            raise HTTPException(status_code=400, detail="The selected worksheet must contain a Property Number (or Assessment Number) and an Owner/Occupier Name column.")
        _, field_meta = column_map(labels)
        imported = 0
        skipped = 0
        errors = []
        lg = None
        if local_government.strip():
            lg = db.query(LocalGovernment).filter(func.lower(LocalGovernment.name) == local_government.strip().lower()).first()
            if lg is None:
                lg = LocalGovernment(name=local_government.strip(), code=(prefix.strip() or None), prefix=(prefix.strip() or None))
                db.add(lg)
                db.flush()

        # Register each detected field; same normalized heading/occurrence reuses the field key.
        fields_by_key = {}
        for meta in field_meta:
            field = db.query(JournalField).filter(JournalField.field_key == meta["field_key"]).first()
            if field is None:
                field = JournalField(**meta)
                db.add(field)
                db.flush()
            else:
                field.label = meta["label"]
                field.position = meta["position"]
                field.occurrence = meta["occurrence"]
                field.normalized_label = meta["normalized_label"]
            fields_by_key[meta["field_key"]] = field

        for excel_row_number, values in enumerate(rows_iter, start=2):
            values = list(values[:len(labels)])
            if not values or all(v is None or str(v).strip() == "" for v in values):
                skipped += 1
                continue
            try:
                property_no = clean_text(values[standard["property_no"]])
                owner_name = clean_text(values[standard["owner_name"]])
                if not property_no or not owner_name:
                    skipped += 1
                    continue
                prop = db.query(Property).filter(Property.property_no == property_no).first()
                if prop is None:
                    prop = Property(property_no=property_no, owner_name=owner_name)
                    db.add(prop)
                    db.flush()
                prop.owner_name = owner_name
                for std in ("lg_code", "address", "rating_area", "annual_value", "rate_due", "year", "status", "lg_prefix"):
                    if std not in standard:
                        continue
                    raw = values[standard[std]] if standard[std] < len(values) else None
                    text_value = clean_text(raw)
                    if std == "lg_code":
                        # Keep an existing LG Code; fill it only when currently blank.
                        if text_value is not None and not prop.lg_code:
                            prop.lg_code = text_value
                    elif std in ("annual_value", "rate_due"):
                        if text_value is not None:
                            setattr(prop, std, numeric_value(raw))
                    elif std == "year":
                        if text_value is not None:
                            n = numeric_value(raw)
                            if n is not None:
                                prop.year = int(n)
                    elif text_value is not None:
                        setattr(prop, std, text_value)
                if lg is not None:
                    prop.lg_id = lg.id
                    if prefix.strip() and not prop.lg_prefix:
                        prop.lg_prefix = prefix.strip()
                db.flush()
                for meta in field_meta:
                    idx = meta["position"]
                    raw = values[idx] if idx < len(values) else None
                    text_value = clean_text(raw)
                    existing = db.query(JournalValue).filter(
                        JournalValue.property_id == prop.id,
                        JournalValue.field_id == fields_by_key[meta["field_key"]].id,
                    ).first()
                    if existing is None:
                        existing = JournalValue(property_id=prop.id, field_id=fields_by_key[meta["field_key"]].id)
                        db.add(existing)
                    # Blank cells don't erase a previously imported value.
                    if text_value is not None:
                        existing.value_text = text_value
                        existing.value_numeric = numeric_value(raw)
                imported += 1
            except Exception as exc:
                errors.append({"row": excel_row_number, "error": str(exc)[:300]})
        db.commit()
        workbook.close()
        return {"success": True, "worksheet": sheet_name, "imported": imported, "skipped": skipped, "errors": errors[:50], "fields_detected": len(field_meta), "message": f"Imported {imported} rows from worksheet '{sheet_name}'."}
    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Excel import failed: {exc}")


@router.get("/fields")
def get_journal_fields(db: Session = Depends(get_db)):
    fields = db.query(JournalField).order_by(JournalField.position, JournalField.id).all()
    return {"fields": [{"field_key": f.field_key, "label": f.label, "position": f.position, "occurrence": f.occurrence} for f in fields]}


def build_query(db: Session, search: str | None, filters: dict[str, str], selected_only: bool = False):
    query = db.query(Property)
    if search and search.strip():
        pattern = f"%{search.strip()}%"
        matching_value_property_ids = db.query(JournalValue.property_id).filter(JournalValue.value_text.ilike(pattern)).subquery()
        query = query.filter(or_(
            Property.property_no.ilike(pattern),
            Property.lg_code.ilike(pattern),
            Property.owner_name.ilike(pattern),
            Property.address.ilike(pattern),
            Property.rating_area.ilike(pattern),
            Property.id.in_(matching_value_property_ids),
        ))
    for key, value in filters.items():
        if not value.strip():
            continue
        field = db.query(JournalField).filter(JournalField.field_key == key).first()
        if field is None:
            continue
        matching_ids = db.query(JournalValue.property_id).filter(
            JournalValue.field_id == field.id,
            JournalValue.value_text.ilike(f"%{value.strip()}%"),
        ).subquery()
        query = query.filter(Property.id.in_(matching_ids))
    return query


@router.get("/records")
def get_journal_records(
    search: str | None = None,
    filters: str = Query("{}", description="JSON object: field_key to contains-text filter"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    sort_by: str = "property_no",
    sort_dir: str = Query("asc", pattern="^(asc|desc)$"),
    db: Session = Depends(get_db),
):
    try:
        filter_map = json.loads(filters or "{}")
        if not isinstance(filter_map, dict):
            raise ValueError("filters must be a JSON object")
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid filters JSON: {exc}")
    query = build_query(db, search, {str(k): str(v) for k, v in filter_map.items()})
    total = query.count()
    if sort_by in {"property_no", "lg_code", "owner_name", "address", "rating_area", "status", "year", "rate_due", "annual_value", "id"}:
        sort_column = getattr(Property, sort_by)
        query = query.order_by(sort_column.desc().nullslast() if sort_dir == "desc" else sort_column.asc().nullslast())
    else:
        field = db.query(JournalField).filter(JournalField.field_key == sort_by).first()
        if field:
            sort_value = db.query(JournalValue.value_text).filter(JournalValue.property_id == Property.id, JournalValue.field_id == field.id).correlate(Property).scalar_subquery()
            query = query.order_by(sort_value.desc().nullslast() if sort_dir == "desc" else sort_value.asc().nullslast())
        else:
            query = query.order_by(Property.property_no.asc())
    items = query.offset((page - 1) * page_size).limit(page_size).all()
    ids = [p.id for p in items]
    dynamic = {}
    if ids:
        rows = db.query(JournalValue.property_id, JournalField.field_key, JournalValue.value_text).join(JournalField, JournalField.id == JournalValue.field_id).filter(JournalValue.property_id.in_(ids)).all()
        for property_id, field_key, value in rows:
            dynamic.setdefault(property_id, {})[field_key] = value
    records = []
    for prop in items:
        rec = safe_record(prop)
        rec["journal_values"] = dynamic.get(prop.id, {})
        records.append(rec)
   
    response_data = {
        "records": records,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": math.ceil(total / page_size) if total else 0,
    }

    return make_json_safe(jsonable_encoder(response_data))


@router.post("/bill-selection")
def get_bill_selection(payload: dict, db: Session = Depends(get_db)):
    """Returns property identifiers matching the same server-side journal filters."""
    search = payload.get("search")
    filters = payload.get("filters") or {}
    if not isinstance(filters, dict):
        raise HTTPException(status_code=400, detail="filters must be an object")
    query = build_query(db, search, {str(k): str(v) for k, v in filters.items()})
    props = query.order_by(Property.property_no.asc()).all()
    return {"count": len(props), "properties": [{"id": p.id, "property_no": p.property_no, "lg_code": p.lg_code, "owner_name": p.owner_name} for p in props]}


@router.get("/export")
def export_journal(
    search: str | None = None,
    filters: str = Query("{}"),
    db: Session = Depends(get_db),
):
    try:
        filter_map = json.loads(filters or "{}")
        if not isinstance(filter_map, dict):
            raise ValueError("filters must be an object")
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid filters JSON: {exc}")
    query = build_query(db, search, {str(k): str(v) for k, v in filter_map.items()})
    props = query.order_by(Property.property_no.asc()).all()
    fields = db.query(JournalField).order_by(JournalField.position, JournalField.id).all()
    values = {}
    if props:
        rows = db.query(JournalValue.property_id, JournalField.field_key, JournalValue.value_text).join(JournalField).filter(JournalValue.property_id.in_([p.id for p in props])).all()
        for pid, key, value in rows:
            values.setdefault(pid, {})[key] = value
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Journal"
    base_headers = ["property_no", "lg_code", "owner_name", "address", "rating_area", "annual_value", "rate_due", "year", "status"]
    ws.append(base_headers + [f.label if f.occurrence == 1 else f"{f.label} ({f.occurrence})" for f in fields])
    for p in props:
        base = [p.property_no, p.lg_code, p.owner_name, p.address, p.rating_area, p.annual_value, p.rate_due, p.year, p.status]
        extra = [values.get(p.id, {}).get(f.field_key) for f in fields]
        ws.append(base + extra)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=master_journal_filtered.xlsx"})
